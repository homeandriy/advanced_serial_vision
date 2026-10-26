from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook

from serial_vision.warehouse.security import WarehouseError

HEADERS = (
    "code",
    "name",
    "address",
    "note",
    "active",
    "responsible_name",
    "position",
    "phone",
    "email",
)


def read_warehouse_xlsx(path: Path) -> list[dict[str, object]]:
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 20 * 1024 * 1024:
        raise WarehouseError("wh_import_invalid")
    try:
        with ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise WarehouseError("wh_import_invalid")
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book["Warehouses"] if "Warehouses" in book.sheetnames else book.active
            if sheet is None:
                raise WarehouseError("wh_import_invalid")
            rows = sheet.iter_rows(values_only=True)
            header = [str(value or "").strip().lower() for value in next(rows)]
            if not all(key in header for key in ("code", "name", "address")):
                raise WarehouseError("wh_import_invalid")
            result = []
            for index, values in enumerate(rows):
                if index >= 10000:
                    raise WarehouseError("wh_import_invalid")
                if not any(value is not None for value in values):
                    continue
                item: dict[str, object] = {
                    key: str(values[header.index(key)] or "")
                    if key in header and header.index(key) < len(values)
                    else ""
                    for key in HEADERS
                }
                active = str(item["active"]).strip().casefold()
                if active not in {"", "true", "1", "yes", "так", "false", "0", "no", "ні"}:
                    raise WarehouseError("wh_import_invalid")
                item["active"] = active not in {"false", "0", "no", "ні"}
                result.append(item)
            if not result:
                raise WarehouseError("wh_import_invalid")
            return result
        finally:
            book.close()
    except (BadZipFile, KeyError, ValueError, StopIteration, OSError) as error:
        raise WarehouseError("wh_import_invalid") from error


def save_template(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Warehouses"
    sheet.append(HEADERS)
    book.save(path)
    book.close()
